#!/usr/bin/env python3
"""Read actual exported documents with an independent spreadsheet implementation."""
import csv, hashlib, io, json, sys
from pathlib import Path
import openpyxl

root=Path(sys.argv[1]);results={}
for database in ('sqlite','postgres'):
    capture=json.loads((root/f'console-{database}.json').read_text())
    report=capture['snapshot']['report']
    assert json.loads((root/f'report-{database}.json').read_text())==report
    wb=openpyxl.load_workbook(root/f'report-{database}.xlsx',read_only=True,data_only=False,keep_links=False)
    assert wb.sheetnames==['概览','主机与基线','漏洞记录','告警处置','每日趋势','漏洞覆盖','来源待处理']
    assert wb['概览']['B2'].value==report['title']
    host_rows=list(wb['主机与基线'].iter_rows(min_row=2,values_only=True))
    assert [r[1] for r in host_rows]==[h['hostname'] for h in report['hosts']]
    assert all(cell.data_type!='f' for sheet in wb for row in sheet for cell in row)
    assert sum(r[2] for r in wb['漏洞记录'].iter_rows(min_row=2,values_only=True))==sum(r['records'] for r in report['vulnerabilities'])
    csv_rows=list(csv.reader(io.StringIO((root/f'report-{database}.csv').read_text(encoding='utf-8-sig'),newline='')))
    hostile=next(h['hostname'] for h in report['hosts'] if h['hostname'].startswith('='))
    assert any("'"+hostile in row for row in csv_rows)
    html=(root/f'report-{database}.html').read_text()
    assert '<img' not in html and '<script' not in html and '&lt;img' in html
    assert '危急（5）' in html and '不代表等保认证结论' in html
    results[database]={'reader':f'openpyxl {openpyxl.__version__}','sheets':wb.sheetnames,'hostRows':len(host_rows),'allTextCellsAreNotFormulas':True,
        'documents':[{'file':f'report-{database}.{ext}','sha256':hashlib.sha256((root/f'report-{database}.{ext}').read_bytes()).hexdigest()} for ext in ('csv','xlsx','html','json')]}
    wb.close()
(root/'document-reader-result.json').write_text(json.dumps({'passed':True,'databases':results},ensure_ascii=False,indent=2)+'\n')
print('Both databases: seven-sheet XLSX, quoted CSV, escaped HTML and full JSON independently verified')
